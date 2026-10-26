"""Build Maryland precinct-to-district crosswalks from official BEFs.

The Redistricting Data Hub 2020 block election file supplies a precinct ID for
each 2020 Census block.  The official Maryland block-equivalency workbooks
assign those same blocks to districts.  We aggregate Census block land area by
precinct and district, falling back to block counts for the handful of
precincts (if any) whose blocks have no land area.
"""

from __future__ import annotations

import csv
import io
import re
import sys
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
VENDOR = ROOT / ".vendor" / "pyshp"
if VENDOR.exists():
    sys.path.insert(0, str(VENDOR))

import shapefile  # type: ignore  # provided in .vendor/pyshp


def reader_from_zip(path: Path) -> shapefile.Reader:
    with zipfile.ZipFile(path) as archive:
        stem = next(name[:-4] for name in archive.namelist() if name.lower().endswith(".shp"))
        return shapefile.Reader(
            shp=io.BytesIO(archive.read(f"{stem}.shp")),
            shx=io.BytesIO(archive.read(f"{stem}.shx")),
            dbf=io.BytesIO(archive.read(f"{stem}.dbf")),
        )


def read_precinct_ids(path: Path) -> dict[str, str]:
    reader = reader_from_zip(path)
    result: dict[str, str] = {}
    for record in reader.iterRecords():
        row = record.as_dict()
        block = str(row.get("GEOID20") or "").strip()
        precinct = " ".join(str(row.get("PRECINCTID") or "").split()).upper()
        if block and precinct:
            result[block] = precinct
    reader.close()
    return result


def read_block_land_area(path: Path) -> dict[str, int]:
    reader = reader_from_zip(path)
    result = {
        str(record["GEOID20"]).strip(): int(record["ALAND20"] or 0)
        for record in reader.iterRecords()
    }
    reader.close()
    return result


def normalize_district(value: object, scope: str) -> str:
    district = str(value or "").strip().upper()
    if scope == "state_senate" and district.startswith("SD"):
        district = district[2:]
    if scope == "state_house" and district.startswith("HD"):
        district = district[2:]
    match = re.fullmatch(r"0*(\d+)([A-Z]*)", district)
    if match:
        return f"{int(match.group(1))}{match.group(2)}"
    return district


def read_bef(path: Path, scope: str) -> dict[str, str]:
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header or str(header[0]).strip().lower() != "block":
        raise ValueError(f"Unexpected BEF header in {path.name}: {header!r}")

    result: dict[str, str] = {}
    for row in rows:
        block = str(row[0] or "").strip()
        district = normalize_district(row[1], scope) if len(row) > 1 else ""
        if block and district:
            result[block] = district
    book.close()
    return result


def build_rows(
    precinct_by_block: dict[str, str],
    land_area_by_block: dict[str, int],
    district_by_block: dict[str, str],
) -> tuple[list[dict[str, object]], dict[str, int]]:
    pair_area: dict[tuple[str, str], int] = defaultdict(int)
    pair_blocks: dict[tuple[str, str], int] = defaultdict(int)
    precinct_area: dict[str, int] = defaultdict(int)
    precinct_blocks: dict[str, int] = defaultdict(int)
    missing_precinct = 0

    for block, district in district_by_block.items():
        precinct = precinct_by_block.get(block)
        if not precinct:
            missing_precinct += 1
            continue
        land_area = land_area_by_block.get(block, 0)
        key = (precinct, district)
        pair_area[key] += land_area
        pair_blocks[key] += 1
        precinct_area[precinct] += land_area
        precinct_blocks[precinct] += 1

    rows: list[dict[str, object]] = []
    for precinct, district in sorted(pair_blocks):
        land_area = pair_area[(precinct, district)]
        block_count = pair_blocks[(precinct, district)]
        total_land_area = precinct_area[precinct]
        total_blocks = precinct_blocks[precinct]
        weight = land_area / total_land_area if total_land_area else block_count / total_blocks
        rows.append(
            {
                "precinct_key": precinct,
                "district_num": district,
                "area_weight": f"{weight:.12f}",
                "land_area": land_area,
                "precinct_land_area": total_land_area,
                "block_count": block_count,
                "precinct_block_count": total_blocks,
            }
        )

    diagnostics = {
        "bef_blocks": len(district_by_block),
        "matched_blocks": sum(pair_blocks.values()),
        "missing_precinct_blocks": missing_precinct,
        "precincts": len(precinct_blocks),
        "rows": len(rows),
    }
    return rows, diagnostics


def write_crosswalk(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "precinct_key",
        "district_num",
        "area_weight",
        "land_area",
        "precinct_land_area",
        "block_count",
        "precinct_block_count",
    ]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    precinct_by_block = read_precinct_ids(DATA / "md_2020_gen_2020_blocks.zip")
    land_area_by_block = read_block_land_area(DATA / "tl_2020_24_tabblock20.zip")
    output_dir = DATA / "crosswalks_bef"
    plans = [
        (
            "precinct_to_congressional_2022_enacted.csv",
            "congressional",
            DATA / "SB1012-2022-Md-Congress-BEQ.xlsx",
        ),
        (
            "precinct_to_congressional_2021_mcrc.csv",
            "congressional",
            DATA / "MCRC-2021-Final-Congressional-BEF.xlsx",
        ),
        (
            "precinct_to_state_house_2022.csv",
            "state_house",
            DATA / "MD-2022OFFICIALLYADOPTEDHOUSE.xlsx",
        ),
        (
            "precinct_to_state_senate_2022.csv",
            "state_senate",
            DATA / "MD-2022OFFICIALLYADOPTEDSENATE.xlsx",
        ),
    ]

    for output_name, scope, source in plans:
        assignments = read_bef(source, scope)
        rows, diagnostics = build_rows(precinct_by_block, land_area_by_block, assignments)
        output = output_dir / output_name
        write_crosswalk(output, rows)
        print(
            f"wrote {output.relative_to(ROOT)}: {diagnostics['rows']} rows, "
            f"{diagnostics['precincts']} precincts, {diagnostics['matched_blocks']}/"
            f"{diagnostics['bef_blocks']} blocks matched"
        )
        if diagnostics["missing_precinct_blocks"]:
            raise RuntimeError(
                f"{source.name}: {diagnostics['missing_precinct_blocks']} BEF blocks lack precinct IDs"
            )


if __name__ == "__main__":
    main()
